import os
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

W, H = 1920, 1080
slides_dir = r'c:\Users\prajw\Desktop\RedFlow\docs\slides'
cache_dir = os.path.join(slides_dir, 'cache_40')
os.makedirs(cache_dir, exist_ok=True)

logo_path = r'c:\Users\prajw\Desktop\RedFlow\frontend\src\assets\logo.png'
logo_orig = Image.open(logo_path).convert('RGBA')

def get_font(size, bold=True):
    font_names = ['segoeuib.ttf' if bold else 'segoeui.ttf',
                  'arialbd.ttf' if bold else 'arial.ttf',
                  'calibrib.ttf' if bold else 'calibri.ttf']
    for fn in font_names:
        p = os.path.join(os.environ.get('WINDIR', 'C:\\Windows'), 'Fonts', fn)
        if os.path.exists(p):
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()

font_brand_bold = get_font(34, bold=True)
font_num_bold = get_font(24, bold=True)
font_footer = get_font(14, bold=False)

print('Rendering base template with 14% watermark...')

# 1. Render Base Background (without page number)
base_img = Image.new('RGBA', (W, H))
draw = ImageDraw.Draw(base_img)

# Light gradient
for y in range(H):
    r = int(255 - (y / H) * 6)
    g = int(255 - (y / H) * 6)
    b = int(255 - (y / H) * 5)
    draw.line([(0, y), (W, y)], fill=(r, g, b, 255))

# Top gradient bar (Crimson to Navy)
accent_h = 5
for x in range(W):
    ratio = x / W
    r = int(220 * (1 - ratio) + 23 * ratio)
    g = int(38 * (1 - ratio) + 43 * ratio)
    b = int(38 * (1 - ratio) + 77 * ratio)
    draw.line([(x, 0), (x, accent_h - 1)], fill=(r, g, b, 255))

# Watermark logo (14% prominent)
wm_size = 680
wm_logo = logo_orig.resize((wm_size, wm_size), Image.Resampling.LANCZOS)
wm_data = list(wm_logo.getdata())
adjusted_wm = []
watermark_opacity = 0.14
for item in wm_data:
    adjusted_wm.append((item[0], item[1], item[2], int(item[3] * watermark_opacity)))
wm_logo.putdata(adjusted_wm)

wm_x = (W - wm_size) // 2
wm_y = (H - wm_size) // 2 + 10
base_img.alpha_composite(wm_logo, (wm_x, wm_y))

# Top-Right Branding: Logo + 'Red' + 'Flow'
brand_logo_size = 36
small_logo = logo_orig.resize((brand_logo_size, brand_logo_size), Image.Resampling.LANCZOS)

red_bbox = draw.textbbox((0, 0), 'Red', font=font_brand_bold)
red_w = red_bbox[2] - red_bbox[0]
flow_bbox = draw.textbbox((0, 0), 'Flow', font=font_brand_bold)
flow_w = flow_bbox[2] - flow_bbox[0]

margin_right = 72
top_y = 44
gap = 10

total_brand_w = brand_logo_size + gap + red_w + flow_w
start_x = W - margin_right - total_brand_w

base_img.alpha_composite(small_logo, (start_x, top_y))
text_x = start_x + brand_logo_size + gap

draw.text((text_x, top_y + 1), 'Red', font=font_brand_bold, fill=(211, 47, 47, 255))
draw.text((text_x + red_w, top_y + 1), 'Flow', font=font_brand_bold, fill=(23, 43, 77, 255))

# Footer separator line
footer_y = H - 64
draw.line([(72, footer_y), (W - 72, footer_y)], fill=(226, 232, 240, 255), width=1)

# Footer Left text
draw.text((72, footer_y + 16), 'RedFlow  •  Decentralized Blood Supply & Traceability Platform', font=font_footer, fill=(148, 163, 184, 255))

print('Generating 40 slides and PPTX presentation...')

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6] # Blank slide layout

for i in range(1, 41):
    page_str = f"{i:02d}"
    slide_img = base_img.copy()
    s_draw = ImageDraw.Draw(slide_img)
    
    # Page number bottom right
    num_bbox = s_draw.textbbox((0, 0), page_str, font=font_num_bold)
    nw = num_bbox[2] - num_bbox[0]
    nh = num_bbox[3] - num_bbox[1]
    
    btn_w = max(52, nw + 26)
    btn_h = 32
    bx1 = W - 72 - btn_w
    by1 = footer_y + 11
    bx2 = bx1 + btn_w
    by2 = by1 + btn_h
    
    # Rounded badge
    s_draw.rounded_rectangle([bx1, by1, bx2, by2], radius=6, fill=(241, 245, 249, 255), outline=(226, 232, 240, 255), width=1)
    tx = bx1 + (btn_w - nw) // 2
    ty = by1 + (btn_h - nh) // 2 - num_bbox[1]
    s_draw.text((tx, ty), page_str, font=font_num_bold, fill=(211, 47, 47, 255))
    
    img_path = os.path.join(cache_dir, f'slide_{page_str}.png')
    slide_img.save(img_path, 'PNG')
    
    # Add slide
    slide = prs.slides.add_slide(blank_layout)
    
    # Add picture directly to the slide shapes (full bleed 13.333" x 7.5")
    pic = slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    
    # Lock the picture so user cannot accidentally drag, rotate, or resize it
    cNvPicPr = pic._element.nvPicPr.find('{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPicPr')
    if cNvPicPr is not None:
        picLocks = cNvPicPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}picLocks')
        if picLocks is not None:
            picLocks.set('noMove', '1')
            picLocks.set('noResize', '1')
            picLocks.set('noRot', '1')
            picLocks.set('noChangeAspect', '1')

# Save presentation to multiple targets so if one is locked by PowerPoint, the others succeed
target_paths = [
    r'c:\Users\prajw\Desktop\RedFlow\RedFlow_Presentation_40_Pages.pptx',
    r'c:\Users\prajw\Desktop\RedFlow\docs\slides\RedFlow_Presentation_40_Pages.pptx',
    os.path.join(r'C:\Users\prajw\.gemini\antigravity\brain\640542b5-c6cd-44d9-adaf-6a1e585a7de3', 'RedFlow_Presentation_40_Pages.pptx'),
    r'c:\Users\prajw\Desktop\RedFlow\RedFlow_Presentation_40_Slides.pptx',
    r'c:\Users\prajw\Desktop\RedFlow\docs\slides\RedFlow_Presentation_40_Slides.pptx'
]

saved_files = []
for p in target_paths:
    try:
        prs.save(p)
        saved_files.append(p)
        print(f'Successfully saved: {p}')
    except Exception as e:
        print(f'Skipped locked file {p}: {e}')

print('Done!')
